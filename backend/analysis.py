"""Dividend and covered-call analysis. Refactored from notebooks/DividendvsCovercalls.ipynb
so the same logic powers the website API, the Excel report and the batch CLI."""
import io, json, os, re, time
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime

import pandas as pd
import yfinance as yf
from openpyxl.styles import PatternFill

HORIZONS = {"weekly": 7, "monthly": 30, "quarterly": 90}   # target days to expiration
TICKER_RE = re.compile(r"^[A-Z][A-Z0-9.\-]{0,9}$")
EXCLUDE = {t.strip().upper() for t in os.getenv("EXCLUDE_TICKERS", "").split(",") if t.strip()}
MAX_TICKERS = int(os.getenv("MAX_TICKERS", "60"))
FILLS = {k: PatternFill("solid", start_color=c, end_color=c) for k, c in {0: "C6EFCE", 1: "FFC000", 2: "FF6666"}.items()}
CALL_COLS = ["Institution", "Ticker", "Expiration", "Days", "Contracts", "Price", "Strike", "OTM %", "Premium / Share",
             "Bid", "Ask", "Volume", "Open Interest", "IV %", "Total Premium", "Return %", "Annualized %",
             "Breakeven", "Dividend Risk", "Earnings Risk", "Comments", "Group"]
_cache = {}


def _stock(ticker):
    """(yf.Ticker, info) with a 5-minute cache so repeat lookups are fast and gentle on the data provider."""
    hit = _cache.get(ticker)
    if hit and time.time() - hit[0] < 300:
        return hit[1], hit[2]
    if len(_cache) > 500:
        _cache.clear()
    stock = yf.Ticker(ticker)
    info = stock.info or {}
    _cache[ticker] = (time.time(), stock, info)
    return stock, info


def _to_dt(v):
    if v in (None, "", 0):
        return None
    try:
        if isinstance(v, (int, float)):
            return datetime.fromtimestamp(v)
        t = pd.Timestamp(v)
        if pd.isna(t):
            return None
        return (t.tz_localize(None) if t.tzinfo else t).to_pydatetime()
    except Exception:
        return None


def _price(info):
    return info.get("currentPrice") or info.get("regularMarketPrice") or info.get("previousClose")


def events(stock, info, expiry):
    """Dividend/earnings risk between today and the option expiry."""
    today, notes, div, earn = date.today(), [], False, False
    ex = _to_dt(info.get("exDividendDate"))
    if ex and today <= ex.date() <= expiry:
        div = True
        notes.append("Ex-dividend risk")
    try:
        cal, dates = stock.calendar, []
        if isinstance(cal, dict):                       # newer yfinance returns a dict
            dates = cal.get("Earnings Date") or []
        elif cal is not None and not cal.empty and "Earnings Date" in cal.index:
            dates = list(cal.loc["Earnings Date"])      # older versions return a DataFrame
        if not isinstance(dates, (list, tuple)):
            dates = [dates]
        for d in dates:
            t = _to_dt(d)
            if t and today <= t.date() <= expiry:
                earn = True
                notes.append("Earnings announcement")
                break
    except Exception:
        pass
    return div, earn, "; ".join(notes)


def dividend_row(info, ticker, qty, inst):
    price, annual = _price(info), info.get("dividendRate") or 0
    if annual <= 0 or not price:
        return None
    q = round(annual / 4, 4)
    return {"Institution": inst, "Ticker": ticker, "Quantity": qty, "Price": price,
            "Annual Dividend / Share": annual, "Quarterly Dividend / Share": q,
            "Quarterly Dividend Income": round(q * qty, 2), "Annual Dividend Income": round(annual * qty, 2),
            "Dividend Yield %": round(annual / price * 100, 2),
            "Ex-Dividend Date": _to_dt(info.get("exDividendDate")), "Dividend Date": _to_dt(info.get("dividendDate"))}


def covered_calls(stock, info, ticker, qty, inst, horizon, preview=False):
    """Top 5 out-of-the-money calls (by return) for the expiry closest to the horizon."""
    contracts = qty // 100
    if ticker in EXCLUDE or (contracts == 0 and not preview):
        return None
    contracts = max(contracts, 1)
    price = _price(info)
    try:
        exps = stock.options
    except Exception:
        return None
    today = date.today()
    cands = []
    for e in exps or []:
        days = (datetime.strptime(e, "%Y-%m-%d").date() - today).days
        if days >= 1:
            cands.append((abs(days - HORIZONS[horizon]), e, days))
    if not price or not cands:
        return None
    _, expiry, days = min(cands)
    try:
        ch = stock.option_chain(expiry).calls
    except Exception:
        return None
    ch = ch[(ch["strike"] > price) & (ch["volume"] > 10) & (ch["openInterest"] > 50) & (ch["lastPrice"] > 0)]
    if ch.empty:
        return None
    div, earn, comment = events(stock, info, datetime.strptime(expiry, "%Y-%m-%d").date())
    d = ch.rename(columns={"strike": "Strike", "bid": "Bid", "ask": "Ask", "volume": "Volume",
                           "openInterest": "Open Interest"}).copy()
    d["Institution"], d["Ticker"], d["Expiration"], d["Days"] = inst, ticker, expiry, days
    d["Contracts"], d["Price"] = contracts, price
    d["OTM %"] = (d["Strike"] / price - 1) * 100
    d["Premium / Share"] = d["lastPrice"]
    d["IV %"] = d["impliedVolatility"] * 100
    d["Total Premium"] = d["lastPrice"] * 100 * contracts
    d["Return %"] = d["lastPrice"] / price * 100
    d["Annualized %"] = d["Return %"] / days * 365
    d["Breakeven"] = price - d["lastPrice"]
    d["Dividend Risk"], d["Earnings Risk"], d["Comments"] = div, earn, comment
    d = d.sort_values("Return %", ascending=False).head(5)
    d["Group"] = range(1, len(d) + 1)
    d = d[CALL_COLS].copy()
    num = d.select_dtypes("number").columns
    d[num] = d[num].round(2)
    return d


# ---------------- single stock ---------------- #
def analyze_single(ticker, qty, horizon):
    stock, info = _stock(ticker)
    if not _price(info):
        raise LookupError("No market data found for " + ticker)
    c = records(covered_calls(stock, info, ticker, qty, "", horizon, preview=True))
    div = dividend_row(info, ticker, qty, "")
    return {"ticker": ticker, "name": info.get("shortName") or info.get("longName") or ticker,
            "price": _price(info), "shares": qty, "horizon": horizon,
            "dividend": json.loads(json.dumps(div, default=str)) if div else None,
            "calls": c, "expiry": c[0]["Expiration"] if c else None, "days": c[0]["Days"] if c else None,
            "contracts": qty // 100, "shares_short": max(0, 100 - qty) if qty < 100 else 0}


# ---------------- portfolio ---------------- #
def normalize_input(df):
    """Accepts the original input sheet (Ticker, Quantity, Institution) with flexible header names."""
    names = {c: str(c).strip().lower() for c in df.columns}
    pick = lambda keys: next((c for c, n in names.items() if n in keys), None)
    tc, qc, ic = pick({"ticker", "symbol"}), pick({"quantity", "shares", "qty"}), pick({"institution", "account", "broker"})
    if not tc or not qc:
        raise ValueError("Input needs 'Ticker' and 'Quantity' columns (optional: 'Institution').")
    rows, skipped = [], []
    for _, r in df.iterrows():
        t = str(r[tc]).strip().upper()
        try:
            q = int(float(r[qc]))
        except (TypeError, ValueError):
            q = 0
        inst = str(r[ic]).strip() if ic and pd.notna(r[ic]) else "Unknown"
        if TICKER_RE.match(t) and q > 0:
            rows.append((t, q, inst))
        elif t and t != "NAN":
            skipped.append(t + ": invalid ticker or quantity")
    if len(rows) > MAX_TICKERS:
        raise ValueError("Please limit uploads to %d rows." % MAX_TICKERS)
    if not rows:
        raise ValueError("No valid rows found in the file.")
    return rows, skipped


def build_pivot(calls):
    if calls.empty:
        return pd.DataFrame()
    p = pd.pivot_table(calls, values="Total Premium", index=["Institution", "Ticker", "Strike"], columns=["Group"],
                       aggfunc="sum", fill_value=0)
    blocks = []
    for inst in p.index.get_level_values("Institution").unique():
        blk = p[p.index.get_level_values("Institution") == inst]
        tot = blk.sum()
        tot.name = (inst, "TOTAL", "")
        blocks.append(pd.concat([blk, pd.DataFrame([tot])]))
    return pd.concat(blocks)


def analyze_portfolio(df, horizon="weekly"):
    rows, skipped = normalize_input(df)

    def work(r):
        t, q, i = r
        try:
            stock, info = _stock(t)
            return dividend_row(info, t, q, i), covered_calls(stock, info, t, q, i, horizon), None
        except Exception as e:
            return None, None, "%s: data unavailable (%s)" % (t, e.__class__.__name__)

    with ThreadPoolExecutor(max_workers=6) as ex:
        out = list(ex.map(work, rows))
    skipped += [o[2] for o in out if o[2]]
    divs = pd.DataFrame([o[0] for o in out if o[0]])
    cl = [o[1] for o in out if o[1] is not None]
    calls = pd.concat(cl, ignore_index=True) if cl else pd.DataFrame(columns=CALL_COLS)
    monthly = pd.DataFrame()
    if not divs.empty:
        divs["Month"] = divs["Dividend Date"].dt.strftime("%Y-%m").fillna("Unknown")
        monthly = divs.groupby(["Institution", "Ticker", "Month"])["Quarterly Dividend Income"].sum().reset_index()
        monthly.columns = ["Institution", "Ticker", "Month", "Total Dividend Income"]
        tot = monthly.groupby("Institution")["Total Dividend Income"].sum().rename("Institution Total")
        monthly = monthly.merge(tot, on="Institution", how="left")
        divs = divs.drop(columns=["Month"])
    return {"dividends": divs, "monthly": monthly, "calls": calls, "pivot": build_pivot(calls),
            "skipped": skipped, "horizon": horizon}


def records(df):
    if df is None or df.empty:
        return []
    return json.loads(df.reset_index(drop=True).to_json(orient="records", date_format="iso"))


def summary(res):
    d, c = res["dividends"], res["calls"]
    best = c[c["Group"] == 1] if not c.empty else c
    return {"holdings_with_dividends": int(len(d)),
            "quarterly_dividend_income": round(float(d["Quarterly Dividend Income"].sum()), 2) if not d.empty else 0,
            "tickers_with_calls": int(best["Ticker"].nunique()) if not best.empty else 0,
            "top_pick_premium": round(float(best["Total Premium"].sum()), 2) if not best.empty else 0}


def to_json(res):
    return {"horizon": res["horizon"], "summary": summary(res), "dividends": records(res["dividends"]),
            "monthly": records(res["monthly"]), "calls": records(res["calls"]), "skipped": res["skipped"]}


# ---------------- Excel report (same sheets and colours as the notebook) ---------------- #
def to_workbook(res):
    buf = io.BytesIO()
    d, m, c, p = res["dividends"], res["monthly"], res["calls"], res["pivot"]
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        pd.DataFrame({"Item": ["Expiration horizon", "Holdings with dividends", "Tickers with call ideas", "Notes"],
                      "Value": [res["horizon"], len(d), c["Ticker"].nunique() if not c.empty else 0,
                                "; ".join(res["skipped"]) or "none"]}).to_excel(w, sheet_name="Summary", index=False)
        if not d.empty:
            d.to_excel(w, sheet_name="Dividends", index=False)
            m.to_excel(w, sheet_name="Dividend Pivot", index=False)
        if not c.empty:
            c.to_excel(w, sheet_name="Covered Calls", index=False)
            p.to_excel(w, sheet_name="Covered Calls Pivot")
            ws = w.book["Covered Calls"]
            hdr = {x.value: x.column for x in ws[1]}      # find columns by header name, not letter
            for r in range(2, ws.max_row + 1):
                n = int(bool(ws.cell(r, hdr["Dividend Risk"]).value)) + int(bool(ws.cell(r, hdr["Earnings Risk"]).value))
                for col in range(1, ws.max_column + 1):
                    ws.cell(r, col).fill = FILLS[n]
            risk = c.groupby(["Institution", "Ticker"])[["Dividend Risk", "Earnings Risk"]].max()
            wp = w.book["Covered Calls Pivot"]
            for r in range(2, wp.max_row + 1):
                key = (wp.cell(r, 1).value, wp.cell(r, 2).value)
                if key in risk.index:
                    for col in range(1, wp.max_column + 1):
                        wp.cell(r, col).fill = FILLS[int(risk.loc[key].astype(bool).sum())]
    return buf.getvalue()
